"""Prepara os dados reais do Censo 2022 (IBGE) para o painel de Oriximiná.

Lê os arquivos grandes do IBGE uma vez e gera dois arquivos pequenos
que o painel.py usa:

- dados_setores_oriximina.csv: moradores e faixas de idade de cada setor censitário
- dados_bairros_oriximina.csv: os mesmos dados somados por bairro do mapa da Prefeitura

Rode de novo só se trocar os arquivos do IBGE:  python preparar_dados.py
"""
import os

import geopandas as gpd
import pandas as pd

# Os arquivos lidos e gerados ficam na pasta "dados" do painel, qualquer que seja
# a pasta de onde o script for rodado
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "dados"))

PASTA_IBGE = r"C:\Users\Usuário\Downloads\Pasta Mapa de Painel Oriximiná"
ARQ_BASICO = os.path.join(PASTA_IBGE, "Agregados_por_setores_basico_BR_20260520",
                          "Agregados_por_setores_basico_BR.csv")
ARQ_DEMOGRAFIA = os.path.join(PASTA_IBGE, "Agregados_por_setores_demografia_BR",
                              "Agregados_por_setores_demografia_BR.xlsx")
ARQ_PIRAMIDE = os.path.join(PASTA_IBGE, "Tabelas_panorama",
                            "Censo 2022 - Pirâmide etária - Oriximiná (PA).xlsx")
ARQ_CNEFE = os.path.join(PASTA_IBGE, "CNEFE", "1505304_ORIXIMINA.csv")
CODIGO_ORIXIMINA = "1505304"
TOLERANCIA_BORDA = 100   # metros
AREA_FORA_SEDE = "Sede: áreas fora dos bairros mapeados"
AREA_TROMBETAS = "Porto Trombetas"

# Faixas de idade do IBGE por setor (dicionário de dados, "Dicionário não PCT")
FAIXAS_IBGE = {
    "V01031": "0 a 4", "V01032": "5 a 9", "V01033": "10 a 14",
    "V01034": "15 a 19", "V01035": "20 a 24", "V01036": "25 a 29", "V01037": "30 a 39",
    "V01038": "40 a 49", "V01039": "50 a 59",
    "V01040": "60 a 69", "V01041": "70 ou mais",
}
GRUPOS = {
    "0 a 14 anos": ["V01031", "V01032", "V01033"],
    "15 a 59 anos": ["V01034", "V01035", "V01036", "V01037", "V01038", "V01039"],
    "60 anos ou mais": ["V01040", "V01041"],
}


def ler_basico():
    # Total de pessoas (V0001) e situação urbana/rural de cada setor
    df = pd.read_csv(ARQ_BASICO, sep=";", dtype=str, encoding="latin-1",
                     usecols=["CD_SETOR", "SITUACAO", "CD_MUN", "v0001"])
    df = df[df["CD_MUN"] == CODIGO_ORIXIMINA].copy()
    df["Habitantes"] = pd.to_numeric(df["v0001"]).astype(int)
    df["SITUACAO"] = df["SITUACAO"].fillna("Sem moradores")
    return df[["CD_SETOR", "SITUACAO", "Habitantes"]]


def ler_demografia():
    # O arquivo tem o Brasil inteiro (85 MB); lê linha a linha e guarda só Oriximiná
    import openpyxl
    livro = openpyxl.load_workbook(ARQ_DEMOGRAFIA, read_only=True)
    linhas = livro[livro.sheetnames[0]].iter_rows(values_only=True)
    cabecalho = next(linhas)
    dados = [linha for linha in linhas if str(linha[0]).startswith(CODIGO_ORIXIMINA)]
    df = pd.DataFrame(dados, columns=cabecalho).rename(columns={"CD_setor": "CD_SETOR"})
    df["CD_SETOR"] = df["CD_SETOR"].astype(str)
    return df[["CD_SETOR"] + list(FAIXAS_IBGE)]


def ler_totais_oficiais():
    """Total do município em cada faixa de idade do setor, pela Pirâmide etária oficial do IBGE."""
    piramide = pd.read_excel(ARQ_PIRAMIDE)
    piramide["pessoas"] = piramide.iloc[:, 1] + piramide.iloc[:, 2]   # mulheres + homens
    inicio = piramide["Grupo de idade"].str.extract(r"^(\d+)")[0].astype(int)
    limites = {"V01031": (0, 4), "V01032": (5, 9), "V01033": (10, 14), "V01034": (15, 19),
               "V01035": (20, 24), "V01036": (25, 29), "V01037": (30, 39), "V01038": (40, 49),
               "V01039": (50, 59), "V01040": (60, 69), "V01041": (70, 200)}
    return pd.Series({cod: piramide.loc[inicio.between(a, b), "pessoas"].sum()
                      for cod, (a, b) in limites.items()})


def completar_sigilo(df):
    """O IBGE troca por "X" as faixas com poucas pessoas (sigilo).
    Duas coisas são conhecidas: o total de moradores de cada setor e o total de cada
    faixa no município inteiro (Pirâmide etária). As faixas escondidas são preenchidas
    de forma que as duas somas fechem ao mesmo tempo (ajuste proporcional iterativo)."""
    faixas = list(FAIXAS_IBGE)
    valores = df[faixas].apply(pd.to_numeric, errors="coerce")
    escondido = valores.isna()
    if not escondido.any().any():
        df[faixas] = valores.astype(int)
        return df

    falta_setor = (df["Habitantes"] - valores.sum(axis=1)).clip(lower=0)
    falta_faixa = (ler_totais_oficiais() - valores.sum()).clip(lower=0)

    # Começa pela proporção de cada faixa no município e ajusta linhas e colunas até fechar
    estimativa = escondido.astype(float).mul(valores.sum() + 1, axis=1)
    for _ in range(500):
        estimativa = estimativa.mul(falta_setor / estimativa.sum(axis=1).replace(0, 1), axis=0)
        estimativa = estimativa.mul(falta_faixa / estimativa.sum().replace(0, 1), axis=1)

    # Números inteiros: arredonda cada setor mantendo o total dele (maiores restos)
    inteiros = estimativa.apply(lambda s: s.astype(int), axis=0)
    for i in estimativa.index[escondido.any(axis=1)]:
        sobra = int(falta_setor[i] - inteiros.loc[i].sum())
        restos = (estimativa.loc[i] - inteiros.loc[i])[escondido.loc[i]]
        for faixa in restos.sort_values(ascending=False).index[:max(sobra, 0)]:
            inteiros.at[i, faixa] += 1

    # O arredondamento pode deixar uma faixa com 1 ou 2 pessoas a mais e outra a menos
    # no município; troca essas unidades dentro de setores que têm as duas faixas escondidas
    for _ in range(1000):
        diferenca = inteiros.sum() - falta_faixa
        if (diferenca == 0).all():
            break
        sobra_em, falta_em = diferenca.idxmax(), diferenca.idxmin()
        candidatos = inteiros.index[escondido[sobra_em] & escondido[falta_em] & (inteiros[sobra_em] > 0)]
        if candidatos.empty:
            break
        inteiros.at[candidatos[0], sobra_em] -= 1
        inteiros.at[candidatos[0], falta_em] += 1

    df[faixas] = valores.fillna(inteiros).astype(int)
    return df


def ler_domicilios_cnefe():
    """Pontos dos domicílios do CNEFE 2022 (Cadastro Nacional de Endereços do IBGE)."""
    cnefe = pd.read_csv(ARQ_CNEFE, sep=";", dtype=str,
                        usecols=["COD_SETOR", "COD_ESPECIE", "LATITUDE", "LONGITUDE"])
    cnefe = cnefe[cnefe["COD_ESPECIE"].isin(["1", "2"])]   # domicílios particulares e coletivos
    cnefe["CD_SETOR"] = cnefe["COD_SETOR"].str[:15]
    pontos = gpd.GeoDataFrame(
        cnefe[["CD_SETOR"]],
        geometry=gpd.points_from_xy(cnefe["LONGITUDE"].astype(float), cnefe["LATITUDE"].astype(float)),
        crs=4326,
    ).reset_index(drop=True)

    # Parte dos endereços usa códigos de setor que não existem na malha final do Censo;
    # esses são ligados ao setor pela localização do ponto
    mapa_setores = gpd.read_file("oriximina_setores.geojson")[["CD_SETOR", "geometry"]].to_crs(4326)
    sem_setor = ~pontos["CD_SETOR"].isin(mapa_setores["CD_SETOR"])
    achados = gpd.sjoin(pontos[sem_setor].drop(columns="CD_SETOR"), mapa_setores,
                        how="left", predicate="within")
    achados = achados[~achados.index.duplicated()]
    pontos.loc[sem_setor, "CD_SETOR"] = achados["CD_SETOR"]
    return pontos.dropna(subset=["CD_SETOR"])


def distribuir_por_bairro(setores):
    """Reparte os moradores de cada setor entre os bairros conforme onde ficam os
    domicílios dele (CNEFE). Casas a até TOLERANCIA_BORDA metros de um bairro contam
    para ele, porque o mapa dos bairros foi desenhado à mão a partir do PDF do Plano
    Diretor e as bordas não batem perfeitamente com as do IBGE."""
    mapa_bairros = gpd.read_file("oriximina_bairros.geojson")[["nome", "geometry"]].to_crs(31981)
    mapa_setores = gpd.read_file("oriximina_setores.geojson")
    nucleo = mapa_setores.set_index("CD_SETOR")["NM_NU"]

    domicilios = ler_domicilios_cnefe()
    total_setor = domicilios.groupby("CD_SETOR").size()
    pontos = domicilios.to_crs(31981)
    pontos = gpd.sjoin_nearest(pontos, mapa_bairros, how="left", distance_col="distancia")
    pontos = pontos[~pontos.index.duplicated()].merge(setores[["CD_SETOR", "SITUACAO"]], on="CD_SETOR")

    fora = pontos["distancia"] > TOLERANCIA_BORDA
    pontos.loc[fora, "nome"] = None
    # Zona urbana fora dos bairros: Porto Trombetas à parte, o resto é expansão da sede
    urbano_fora = fora & (pontos["SITUACAO"] == "Urbana")
    eh_trombetas = pontos["CD_SETOR"].map(nucleo).eq("Porto Trombetas")
    pontos.loc[urbano_fora & eh_trombetas, "nome"] = AREA_TROMBETAS
    pontos.loc[urbano_fora & ~eh_trombetas, "nome"] = AREA_FORA_SEDE
    pontos = pontos.dropna(subset=["nome"])   # casas da zona rural ficam fora da visão urbana

    contagem = pontos.groupby(["CD_SETOR", "nome"]).size().rename("casas").reset_index()
    contagem["fracao"] = contagem["casas"] / contagem["CD_SETOR"].map(total_setor)

    # Setores urbanos sem nenhum domicílio no CNEFE: reparte pela área, como antes
    sem_pontos = setores[(setores["SITUACAO"] == "Urbana") & ~setores["CD_SETOR"].isin(total_setor.index)]
    if not sem_pontos.empty:
        geo = mapa_setores[["CD_SETOR", "geometry"]].merge(sem_pontos[["CD_SETOR"]]).to_crs(31981)
        geo["area_setor"] = geo.area
        pedacos = gpd.overlay(geo, mapa_bairros, how="intersection")
        pedacos["fracao"] = pedacos.area / pedacos["area_setor"]
        contagem = pd.concat([contagem, pedacos[["CD_SETOR", "nome", "fracao"]]])

    colunas = ["Habitantes"] + list(GRUPOS)
    partes = contagem.merge(setores, on="CD_SETOR")
    partes[colunas] = partes[colunas].mul(partes["fracao"], axis=0)
    bairros = partes.groupby("nome", as_index=False)[colunas].sum()

    # Arredonda sem perder ninguém: "Habitantes" é sempre a soma exata das três faixas
    for grupo in GRUPOS:
        bairros[grupo] = bairros[grupo].round().astype(int)
    bairros["Habitantes"] = bairros[list(GRUPOS)].sum(axis=1)
    return bairros


def main():
    setores = ler_basico().merge(ler_demografia(), on="CD_SETOR", how="left")
    setores[list(FAIXAS_IBGE)] = setores[list(FAIXAS_IBGE)].fillna(0)   # setores sem moradores
    setores = completar_sigilo(setores)
    for grupo, colunas in GRUPOS.items():
        setores[grupo] = setores[colunas].sum(axis=1)
    setores = setores[["CD_SETOR", "SITUACAO", "Habitantes"] + list(GRUPOS)]
    setores.to_csv("dados_setores_oriximina.csv", sep=";", index=False, encoding="utf-8")

    # Bairros: moradores de cada setor repartidos pelos domicílios do CNEFE (SIRGAS 2000 / UTM 21S)
    bairros = distribuir_por_bairro(setores)
    bairros.to_csv("dados_bairros_oriximina.csv", sep=";", index=False, encoding="utf-8")

    print(f"Setores: {len(setores)} | moradores: {setores['Habitantes'].sum()}")
    print(setores.groupby("SITUACAO")["Habitantes"].sum().to_string())
    print(setores[list(GRUPOS)].sum().to_string())
    print(f"\nBairros: {len(bairros)} | moradores na zona urbana: {bairros['Habitantes'].sum()}")
    print(bairros.sort_values("Habitantes", ascending=False).to_string(index=False))


if __name__ == "__main__":
    main()
